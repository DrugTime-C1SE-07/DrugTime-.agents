#!/usr/bin/env python3
"""Dựng bộ test case Excel từ spec JSON của task.

    python .agents/scripts/testcase_xlsx.py validate [<task>]
    python .agents/scripts/testcase_xlsx.py build    [<task>]                 -> work/<task>/testcases.xlsx
    python .agents/scripts/testcase_xlsx.py merge    [<task>] [--into <xlsx>] -> file tổng (vai trò `testcase-book`)

Spec: work/<task>/testcases.json (cấu trúc trong skill wf-testcase).
Mỗi chức năng là một sheet 9 cột (Test Case ID … Comments), có khối tổng Pass/Fail/Pending/Total,
bảng EP và bảng quyết định bên phải. Sheet `Report` (tổng hợp, biểu đồ, Bug Report) luôn được dựng lại
từ mọi sheet chức năng trong workbook.

`build` ghi đè file của task. `merge` thay sheet có cùng mã chức năng trong file tổng, giữ nguyên
sheet chức năng khác, và giữ mô tả, priority, bug của chức năng khác trong Report cũ.
Exit code: 0 = OK, 1 = spec sai hoặc không ghi được file, 2 = dùng sai hoặc thiếu openpyxl.
"""
from __future__ import annotations

import json
import math
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _lib import ROOT, project_roles, resolve_task_dir  # noqa: E402

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, quote_sheetname
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    Workbook = None

SPEC_NAME = "testcases.json"
XLSX_NAME = "testcases.xlsx"
REPORT = "Report"
STATUSES = ("Pass", "Fail", "Pending")
SEVERITIES = ("Critical", "High", "Medium", "Low")
HEADERS = ["Test Case ID", "Description", "Precondition", "Test Steps", "Test Data",
           "Expected Result", "Actual Result", "Status", "Comments(if any)"]
WIDTHS = [18, 40, 32, 50, 32, 45, 45, 12, 34]
INFO = ["Project Name", "Module Name", "Created By", "Creation Date", "Reviewed By", "Reviewed Date"]
HEADER_ROW = 8
FIRST_ROW = 9
SIDE_COL = 11  # cột K: bảng EP
EP_WIDTHS = [22, 34, 7, 34, 7]
DT_COL = 17  # cột Q: bảng quyết định
DT_NAME_WIDTH, DT_RULE_WIDTH = 34, 8
BUG_HEADERS = ["Bug ID", "Description", "Steps to Reproduce", "Input Data", "Expected Result",
               "Actual Result", "Severity", "Name of Tester", "Logs"]
FUNC_HEADERS = ["Function", "Description", "Sheet Name", "% TCs Executed", "% TCs Passed",
                "% TCs Failed", "Priority", "Remarks", "Done by"]

FONT = "Times New Roman"
GREEN = "6AA84F"
LIGHT_GREEN = "A8D08D"
YELLOW = "FFFF00"
MODULE_RE = re.compile(r"\(([A-Z0-9]+)\)\s*$")


# ---------------------------------------------------------------- spec

def load_spec(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise LookupError(f"Không có {path}.")
    except json.JSONDecodeError as e:
        raise LookupError(f"{path.name} không phải JSON hợp lệ: dòng {e.lineno}, cột {e.colno}: {e.msg}")


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _valid_date(v: str) -> bool:
    if not v:
        return True
    try:
        date.fromisoformat(v)
        return True
    except ValueError:
        return False


def validate(spec: dict) -> list[str]:
    """Trả về danh sách lỗi; rỗng nghĩa là spec dùng được."""
    errs: list[str] = []
    funcs = spec.get("functions")
    if not _text(spec.get("project")):
        errs.append("thiếu 'project'")
    if not isinstance(funcs, list) or not funcs:
        return errs + ["'functions' phải là danh sách không rỗng"]
    codes, sheets, ids = set(), set(), set()
    for fi, f in enumerate(funcs):
        where = f"functions[{fi}]"
        code = _text(f.get("code"))
        if not re.fullmatch(r"[A-Z0-9]+", code):
            errs.append(f"{where}: 'code' phải là chữ in hoa hoặc số, ví dụ AUTH")
        elif code in codes:
            errs.append(f"{where}: trùng code {code}")
        codes.add(code)
        sheet = _text(f.get("sheet")) or _text(f.get("name"))
        if not _text(f.get("name")):
            errs.append(f"{where}: thiếu 'name'")
        if len(sheet) > 31 or re.search(r"[\[\]:*?/\\]", sheet):
            errs.append(f"{where}: tên sheet '{sheet}' dài quá 31 ký tự hoặc chứa []:*?/\\")
        elif sheet.lower() == REPORT.lower():
            errs.append(f"{where}: tên sheet '{REPORT}' dành cho báo cáo")
        elif sheet.lower() in sheets:
            errs.append(f"{where}: trùng tên sheet '{sheet}'")
        sheets.add(sheet.lower())
        if not _text(f.get("created_by")):
            errs.append(f"{where}: thiếu 'created_by'")
        for k in ("created_date", "reviewed_date"):
            if not _valid_date(_text(f.get(k))):
                errs.append(f"{where}: '{k}' phải có dạng YYYY-MM-DD")
        if _text(f.get("priority")) and _text(f.get("priority")) not in SEVERITIES:
            errs.append(f"{where}: 'priority' phải là một trong {', '.join(SEVERITIES)}")
        for ei, ep in enumerate(f.get("ep") or []):
            if not _text(ep.get("field")) or not (ep.get("valid") or ep.get("invalid")):
                errs.append(f"{where}.ep[{ei}]: cần 'field' và ít nhất một lớp valid hoặc invalid")
        for di, dt in enumerate(f.get("decision_tables") or []):
            conds, acts = dt.get("conditions") or [], dt.get("actions") or []
            n = len((conds[0].get("values") or [])) if conds else 0
            if not conds or not acts or n == 0:
                errs.append(f"{where}.decision_tables[{di}]: cần 'conditions' (có 'values') và 'actions'")
                continue
            if any(len(c.get("values") or []) != n for c in conds):
                errs.append(f"{where}.decision_tables[{di}]: mọi điều kiện phải có cùng số rule ({n})")
            for a in acts:
                if any(not isinstance(r, int) or not 1 <= r <= n for r in a.get("rules") or []):
                    errs.append(f"{where}.decision_tables[{di}]: action '{_text(a.get('name'))}' có rule ngoài 1..{n}")
        cases = f.get("cases")
        if not isinstance(cases, list) or not cases:
            errs.append(f"{where}: 'cases' phải là danh sách không rỗng")
            continue
        for ci, c in enumerate(cases):
            cid = _text(c.get("id"))
            w = f"{where}.cases[{ci}] ({cid or 'không có id'})"
            if not re.fullmatch(rf"TC-{re.escape(code)}-[A-Z0-9-]*\d+", cid):
                errs.append(f"{w}: id phải có dạng TC-{code}-NN hoặc TC-{code}-<KỸ THUẬT>-NN")
            elif cid in ids:
                errs.append(f"{w}: trùng id")
            ids.add(cid)
            for k in ("description", "expected", "comments"):
                if not _text(c.get(k)):
                    errs.append(f"{w}: thiếu '{k}'")
            steps = c.get("steps")
            if not isinstance(steps, list) or not [s for s in steps if _text(s)]:
                errs.append(f"{w}: 'steps' phải là danh sách bước không rỗng")
            status = _text(c.get("status"))
            actual = _text(c.get("actual"))
            if status not in STATUSES:
                errs.append(f"{w}: 'status' phải là một trong {', '.join(STATUSES)}")
            elif status == "Pending" and actual:
                errs.append(f"{w}: Pending nghĩa là chưa chạy, không được có 'actual'")
            elif status in ("Pass", "Fail") and not actual:
                errs.append(f"{w}: {status} cần 'actual' lấy từ lần chạy thật")
            if status == "Fail":
                if _text(c.get("severity")) not in SEVERITIES:
                    errs.append(f"{w}: Fail cần 'severity' là một trong {', '.join(SEVERITIES)}")
                if not _text(c.get("root_cause")):
                    errs.append(f"{w}: Fail cần 'root_cause' (đi vào cột Logs của Bug Report)")
    return errs


# ---------------------------------------------------------------- style

def _side(style: str) -> Border:
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)


def _style(cell, *, bold=False, size=12, fill=None, h="center", wrap=True, border="medium", color=None):
    cell.font = Font(name=FONT, size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = _side(border)


def _height(values: list[str], widths: list[float]) -> float:
    """Ước lượng chiều cao dòng cho ô wrap (openpyxl không tự co giãn)."""
    lines = 1
    for v, w in zip(values, widths):
        per = max(int(w * 0.95) - 1, 1)  # Times New Roman 12 rộng hơn font mặc định
        n = sum(max(1, math.ceil(len(part) / per)) for part in str(v or "").split("\n"))
        lines = max(lines, n)
    return max(30.0, lines * 16.5 + 8)


def _steps_text(steps) -> str:
    out = []
    for i, s in enumerate([_text(s) for s in steps if _text(s)], 1):
        out.append(s if re.match(r"^\d+\.\s", s) else f"{i}. {s}")
    return "\n".join(out)


def _date(v: str):
    return date.fromisoformat(v) if v else None


# ---------------------------------------------------------------- function sheet

def write_function_sheet(ws, project: str, f: dict) -> None:
    code = _text(f["code"])
    info = [project, f"{_text(f['name'])} ({code})", _text(f.get("created_by")),
            _date(_text(f.get("created_date"))) or date.today(), _text(f.get("reviewed_by")),
            _date(_text(f.get("reviewed_date")))]
    for r, (label, value) in enumerate(zip(INFO, info), 1):
        _style(ws.cell(r, 1, label), bold=True, fill=GREEN)
        c = ws.cell(r, 2, value)
        _style(c, fill=GREEN)
        if isinstance(value, date):
            c.number_format = "dd/mm/yyyy"
    for i, (h, w) in enumerate(zip(HEADERS, WIDTHS), 1):
        _style(ws.cell(HEADER_ROW, i, h), bold=True, fill=GREEN)
        ws.column_dimensions[get_column_letter(i)].width = w

    cases = f["cases"]
    for r, c in enumerate(cases, FIRST_ROW):
        values = [_text(c["id"]), _text(c["description"]), _text(c.get("precondition")),
                  _steps_text(c["steps"]), _text(c.get("data")), _text(c["expected"]),
                  _text(c.get("actual")), _text(c["status"]), _text(c["comments"])]
        for col, v in enumerate(values, 1):
            _style(ws.cell(r, col, v), h="left" if col in (3, 4, 5, 6, 7, 9) else "center")
        ws.row_dimensions[r].height = _height(values, WIDTHS)
    last = FIRST_ROW + len(cases) - 1
    status_rng = f"H{FIRST_ROW}:H{last}"

    red = PatternFill("solid", bgColor="FFC7CE")
    ws.conditional_formatting.add(status_rng, CellIsRule(operator="equal", formula=['"Fail"'],
                                                         fill=red, font=Font(color="9C0006")))
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(status_rng)
    ws.freeze_panes = f"A{FIRST_ROW}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:I{last}"

    s = last + 3
    rows = [("Pass", f'=COUNTIF({status_rng},"Pass")'), ("Fail", f'=COUNTIF({status_rng},"Fail")'),
            ("Pending", f'=COUNTIF({status_rng},"Pending")'), ("Total", f"=SUM(H{s}:H{s + 2})")]
    for i, (label, formula) in enumerate(rows):
        _style(ws.cell(s + i, 7, label), bold=True, fill=LIGHT_GREEN, border="thin", size=11)
        _style(ws.cell(s + i, 8, formula), fill=LIGHT_GREEN, border="thin", size=11)

    _write_side_tables(ws, f)


def _grow(ws, r: int, values: list[str], widths: list[float]) -> None:
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, _height(values, widths))


def _write_side_tables(ws, f: dict) -> None:
    """Bảng EP ở cột K–O, bảng quyết định ở cột Q trở đi (đặt cạnh nhau như mẫu)."""
    ep = f.get("ep") or []
    if ep:
        col, widths = SIDE_COL, EP_WIDTHS
        _style(ws.cell(HEADER_ROW - 1, col, "Equivalence Partitioning"), bold=True, h="left", border=None)
        for i, (h, w) in enumerate(zip(["Field", "Valid", "Num", "Invalid", "Num"], widths)):
            _style(ws.cell(HEADER_ROW, col + i, h), bold=True, fill=GREEN, border="thin")
            ws.column_dimensions[get_column_letter(col + i)].width = w
        n = 1
        valid_nums = []
        for e in ep:  # đánh số lớp hợp lệ trước, rồi tới lớp không hợp lệ (như mẫu)
            valid_nums.append(list(range(n, n + len(e.get("valid") or []))))
            n += len(e.get("valid") or [])
        for i, e in enumerate(ep):
            inv = e.get("invalid") or []
            inv_nums = list(range(n, n + len(inv)))
            n += len(inv)
            r = HEADER_ROW + 1 + i
            vals = [_text(e["field"]), "\n".join(map(_text, e.get("valid") or [])),
                    "\n".join(map(str, valid_nums[i])), "\n".join(map(_text, inv)), "\n".join(map(str, inv_nums))]
            for j, v in enumerate(vals):
                _style(ws.cell(r, col + j, v), h="left", border="thin", bold=(j == 0))
            _grow(ws, r, vals, widths)
    tables = f.get("decision_tables") or []
    if not tables:
        return
    col, row = DT_COL, HEADER_ROW
    rules = max(len(dt["conditions"][0]["values"]) for dt in tables)
    ws.column_dimensions[get_column_letter(col)].width = DT_NAME_WIDTH
    for k in range(rules):
        ws.column_dimensions[get_column_letter(col + 1 + k)].width = DT_RULE_WIDTH
    for dt in tables:
        conds, acts = dt["conditions"], dt["actions"]
        n = len(conds[0]["values"])
        _style(ws.cell(row - 1, col, _text(dt.get("title")) or "Decision Table"), bold=True, h="left", border=None)
        for block, items in (("Rule", conds), ("Action", acts)):
            _style(ws.cell(row, col, block), bold=True, fill=GREEN, border="thin", h="left")
            for k in range(n):
                _style(ws.cell(row, col + 1 + k, f"R{k + 1}"), bold=True, fill=GREEN, border="thin")
            row += 1
            for it in items:
                _style(ws.cell(row, col, _text(it["name"])), h="left", border="thin")
                marks = it.get("values") if block == "Rule" else ["X" if k + 1 in (it.get("rules") or []) else ""
                                                                 for k in range(n)]
                for k, v in enumerate(marks):
                    _style(ws.cell(row, col + 1 + k, _text(v)), border="thin")
                _grow(ws, row, [_text(it["name"])], [DT_NAME_WIDTH])
                row += 1
        row += 3


# ---------------------------------------------------------------- read back

def _is_function_sheet(ws) -> str | None:
    """Mã chức năng nếu sheet có dạng do script này tạo."""
    if ws.title == REPORT or _text(ws.cell(HEADER_ROW, 1).value) != HEADERS[0]:
        return None
    m = MODULE_RE.search(_text(ws.cell(2, 2).value))
    return m.group(1) if m else None


def _read_cases(ws) -> list[dict]:
    out, r = [], FIRST_ROW
    while _text(ws.cell(r, 1).value):
        out.append({"id": _text(ws.cell(r, 1).value), "actual": _text(ws.cell(r, 7).value),
                    "status": _text(ws.cell(r, 8).value)})
        r += 1
    return out


def _summary_cells(ws) -> dict[str, str]:
    found = {}
    for r in range(FIRST_ROW, ws.max_row + 1):
        label = _text(ws.cell(r, 7).value)
        if label in ("Pass", "Fail", "Pending", "Total") and _text(ws.cell(r, 8).value).startswith("="):
            found[label] = f"{quote_sheetname(ws.title)}!$H${r}"
    return found


def _read_old_report(ws) -> tuple[dict, list[list]]:
    """(meta theo tên sheet, các dòng bug) từ Report cũ."""
    meta, bugs = {}, []
    if ws is None:
        return meta, bugs
    mode = None
    for r in range(1, ws.max_row + 1):
        a = _text(ws.cell(r, 1).value)
        if a == FUNC_HEADERS[0] and _text(ws.cell(r, 3).value) == FUNC_HEADERS[2]:
            mode = "func"
            continue
        if a == BUG_HEADERS[0]:
            mode = "bug"
            continue
        if not a:
            if mode == "func":
                mode = None
            continue
        if mode == "func":
            meta[_text(ws.cell(r, 3).value)] = {"description": _text(ws.cell(r, 2).value),
                                                 "priority": _text(ws.cell(r, 7).value),
                                                 "remarks": _text(ws.cell(r, 8).value)}
        elif mode == "bug":
            bugs.append([ws.cell(r, c).value for c in range(1, 10)])
    return meta, bugs


# ---------------------------------------------------------------- report

def _bug_rows(f: dict) -> list[list]:
    rows = []
    for c in f["cases"]:
        if _text(c["status"]) != "Fail":
            continue
        logs = _text(c.get("root_cause"))
        if _text(c.get("evidence")):
            logs += f"\nLog: {_text(c['evidence'])}"
        rows.append(["BUG-" + _text(c["id"])[3:], _text(c["description"]), _steps_text(c["steps"]),
                     _text(c.get("data")), _text(c["expected"]), _text(c.get("actual")),
                     _text(c["severity"]), _text(f.get("created_by")), logs])
    return rows


def write_report(wb, funcs: list[dict], bugs: list[list]) -> None:
    """funcs: theo thứ tự sheet, mỗi phần tử có name, sheet, tester, description, priority, remarks, cells."""
    if REPORT in wb.sheetnames:
        del wb[REPORT]
    ws = wb.create_sheet(REPORT)
    testers = list(dict.fromkeys(f["tester"] for f in funcs)) or ["-"]
    width_cols = [22, 40, 30, 24, 24, 24, 14, 40, 45]
    for i, w in enumerate(width_cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    last_col = 2 + len(testers) + 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(last_col, 7))
    _style(ws.cell(1, 1, "TEST REPORT"), bold=True, size=18, fill=LIGHT_GREEN, border=None)
    for j, t in enumerate(testers + ["Total"]):
        _style(ws.cell(2, 3 + j, t), bold=True, border="thin")

    def per_tester(label: str) -> list[str]:
        out = []
        for t in testers:
            refs = [f["cells"][label] for f in funcs if f["tester"] == t and label in f["cells"]]
            out.append("=" + "+".join(refs) if refs else 0)
        return out

    lines = [("Executed", "Passed", per_tester("Pass")), ("", "Failed", per_tester("Fail")),
             ("", "(Total) Test executed (Passed + Failed)", None),
             ("Pending", "", per_tester("Pending")), ("(Sub-total) Test Planned", "", None)]
    for i, (a, b, vals) in enumerate(lines):
        r = 3 + i
        _style(ws.cell(r, 1, a), bold=True, border="thin")
        _style(ws.cell(r, 2, b), bold=bool(b), border="thin", h="left")
        for j in range(len(testers)):
            col = get_column_letter(3 + j)
            v = vals[j] if vals else (f"={col}3+{col}4" if i == 2 else f"={col}5+{col}6")
            _style(ws.cell(r, 3 + j, v), border="thin")
        first, end = get_column_letter(3), get_column_letter(2 + len(testers))
        _style(ws.cell(r, 3 + len(testers), f"=SUM({first}{r}:{end}{r})"), bold=True, border="thin")
    ws.merge_cells("A3:A5")

    top = 10
    for i, h in enumerate(FUNC_HEADERS, 1):
        _style(ws.cell(top, i, h), bold=True, fill=GREEN, border="thin")
    for k, f in enumerate(funcs):
        r = top + 1 + k
        c = f["cells"]
        tot = c.get("Total", "0")
        values = [f["name"], f["description"], f["sheet"],
                  f"=IFERROR(({tot}-{c.get('Pending', '0')})/{tot},0)", f"=IFERROR({c.get('Pass', '0')}/{tot},0)",
                  f"=IFERROR({c.get('Fail', '0')}/{tot},0)", f["priority"], f["remarks"], f["tester"]]
        for i, v in enumerate(values, 1):
            cell = ws.cell(r, i, v)
            _style(cell, border="thin", h="left" if i in (2, 8) else "center")
            if 4 <= i <= 6:
                cell.number_format = "0%"
        ws.row_dimensions[r].height = _height(values, width_cols)
    end = top + max(len(funcs), 1)
    red = PatternFill("solid", bgColor="FFC7CE")
    ws.conditional_formatting.add(f"F{top + 1}:F{end}", CellIsRule(operator="greaterThan", formula=["0.2"],
                                                                     fill=red, font=Font(color="9C0006")))
    ws.conditional_formatting.add(f"G{top + 1}:G{end}", FormulaRule(
        formula=[f'OR(G{top + 1}="Critical",G{top + 1}="High")'], fill=red, font=Font(color="9C0006")))
    dv = DataValidation(type="list", formula1='"' + ",".join(SEVERITIES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G{top + 1}:G{end}")

    chart_row = end + 3
    if funcs:
        bar = BarChart()
        bar.type, bar.grouping = "col", "clustered"
        bar.title = "Kết quả theo chức năng"
        bar.y_axis.numFmt = "0%"
        bar.y_axis.scaling.min, bar.y_axis.scaling.max = 0, 1
        bar.add_data(Reference(ws, min_col=4, max_col=6, min_row=top, max_row=end), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=end))
        bar.width, bar.height = 24, 9
        bar.x_axis.delete = bar.y_axis.delete = False  # openpyxl 3.1 mặc định ẩn trục
        ws.add_chart(bar, f"A{chart_row}")
        pie = PieChart()  # Passed/Failed của cột Total; Pending xem ở bảng trên
        pie.title = "Tổng Pass / Fail"
        tot_col = 3 + len(testers)
        pie.add_data(Reference(ws, min_col=tot_col, min_row=3, max_row=4))
        pie.series[0].tx = None
        pie.set_categories(Reference(ws, min_col=2, min_row=3, max_row=4))
        pie.width, pie.height = 12, 9
        ws.add_chart(pie, f"F{chart_row}")

    bug_top = chart_row + 20
    _style(ws.cell(bug_top, 1, "Bug Report"), bold=True, size=14, border=None, h="left")
    for i, h in enumerate(BUG_HEADERS, 1):
        _style(ws.cell(bug_top + 1, i, h), bold=True, fill=YELLOW, border="thin", size=11)
    for k, b in enumerate(bugs):
        r = bug_top + 2 + k
        for i, v in enumerate(b, 1):
            _style(ws.cell(r, i, v), border="thin", h="left", size=11)
        ws.row_dimensions[r].height = _height([_text(v) for v in b], width_cols)
    b_end = bug_top + 1 + max(len(bugs), 1)
    rng = f"G{bug_top + 2}:G{b_end}"
    for sev, color, text in (("Critical", "C00000", "FFFFFF"), ("High", "FF0000", "FFFFFF"),
                             ("Medium", "FFE599", "000000")):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{sev}"'],
                                                      fill=PatternFill("solid", bgColor=color),
                                                      font=Font(color=text, bold=True)))
    dv2 = DataValidation(type="list", formula1='"' + ",".join(SEVERITIES) + '"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(rng)


# ---------------------------------------------------------------- build / merge

def _is_blank_sheet(ws) -> bool:
    # Đọc ô bằng ws.cell() sẽ tạo ô mới, nên xét trực tiếp các ô đã có.
    return all(c.value is None for c in ws._cells.values()) and not ws._images and not ws._charts


def apply(spec: dict, target: Path, fresh: bool) -> list[str]:
    """Ghi spec vào target. Trả về cảnh báo (kết quả cũ trong file tổng bị thay)."""
    warnings: list[str] = []
    wb = Workbook() if fresh or not target.exists() else load_workbook(target)
    old_meta, old_bugs = _read_old_report(wb[REPORT] if REPORT in wb.sheetnames else None)
    new_codes = {_text(f["code"]) for f in spec["functions"]}
    for f in spec["functions"]:
        code, sheet = _text(f["code"]), _text(f.get("sheet")) or _text(f["name"])
        existing = next((ws for ws in wb.worksheets if _is_function_sheet(ws) == code), None)
        clash = sheet in wb.sheetnames and wb[sheet] is not existing
        if clash:
            raise ValueError(f"File đích đã có sheet '{sheet}' không thuộc chức năng {code}.")
        index = None
        if existing is not None:
            new_by_id = {_text(c["id"]): c for c in f["cases"]}
            for old in _read_cases(existing):
                new = new_by_id.get(old["id"])
                if old["status"] in ("Pass", "Fail") and (new is None or _text(new.get("status")) != old["status"]):
                    warnings.append(f"{old['id']}: file đích đang là {old['status']}, spec là "
                                    f"{_text(new.get('status')) if new else 'không còn'}")
            index = wb.worksheets.index(existing)
            del wb[existing.title]
        ws = wb.create_sheet(sheet, index)
        write_function_sheet(ws, _text(spec["project"]), f)
    for ws in list(wb.worksheets):
        if ws.title != REPORT and not _is_function_sheet(ws) and _is_blank_sheet(ws):
            wb.remove(ws)

    spec_by_code = {_text(f["code"]): f for f in spec["functions"]}
    funcs, bugs = [], []
    for ws in wb.worksheets:
        code = _is_function_sheet(ws)
        if not code:
            continue
        name = MODULE_RE.sub("", _text(ws.cell(2, 2).value)).strip()
        f = spec_by_code.get(code)
        meta = old_meta.get(ws.title, {})
        funcs.append({"name": name, "sheet": ws.title, "tester": _text(ws.cell(3, 2).value),
                      "description": _text(f.get("description")) if f else meta.get("description", ""),
                      "priority": _text(f.get("priority")) if f else meta.get("priority", ""),
                      "remarks": _text(f.get("remarks")) if f else meta.get("remarks", ""),
                      "cells": _summary_cells(ws)})
        if f:
            bugs += _bug_rows(f)
        else:
            bugs += [b for b in old_bugs if _text(b[0]).startswith(f"BUG-{code}-")]
    known = {_is_function_sheet(ws) for ws in wb.worksheets}
    bugs += [b for b in old_bugs  # bug của sheet không còn nhận diện được: giữ, không tự xóa
             if not any(_text(b[0]).startswith(f"BUG-{c}-") for c in known | new_codes if c)]
    write_report(wb, funcs, bugs)  # Report luôn là sheet cuối
    wb.calculation.fullCalcOnLoad = True
    target.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(target)
    except PermissionError:
        raise ValueError(f"Không ghi được {target}; file đang mở trong Excel? Đóng file rồi chạy lại.")
    return warnings


def summarize(spec: dict) -> str:
    lines = []
    for f in spec["functions"]:
        st = [_text(c["status"]) for c in f["cases"]]
        lines.append(f"  {_text(f['code'])}: {len(st)} test case — Pass {st.count('Pass')}, "
                     f"Fail {st.count('Fail')}, Pending {st.count('Pending')}")
    return "\n".join(lines)


def default_book() -> Path | None:
    vals = project_roles().get("testcase-book") or []
    return ROOT / vals[0] if vals else None


def main(argv: list[str]) -> int:
    args = argv[1:]
    into = None
    if "--into" in args:
        i = args.index("--into")
        into = args[i + 1] if i + 1 < len(args) else ""
        args = args[:i] + args[i + 2:]
    if not args or args[0] not in ("validate", "build", "merge") or len(args) > 2 or into == "":
        print(__doc__)
        return 2
    if Workbook is None:
        print("LỖI: thiếu openpyxl. Cài bằng: python -m pip install openpyxl")
        return 2
    cmd = args[0]
    try:
        task_dir = resolve_task_dir(args[1] if len(args) > 1 else None)
    except LookupError as e:
        print(f"LỖI: {e}")
        return 2
    try:
        spec = load_spec(task_dir / SPEC_NAME)
    except LookupError as e:
        print(f"LỖI: {e}")
        return 1
    errs = validate(spec)
    if errs:
        print(f"Spec {task_dir.name}/{SPEC_NAME} có {len(errs)} lỗi:")
        for e in errs:
            print(f"  - {e}")
        return 1
    print(f"Spec {task_dir.name}/{SPEC_NAME} hợp lệ:\n{summarize(spec)}")
    if cmd == "validate":
        return 0
    if cmd == "build":
        target = task_dir / XLSX_NAME
    else:
        target = Path(into) if into else default_book()
        if target is None:
            print("LỖI: AGENTS.md chưa khai báo vai trò `testcase-book` trong mục 'Vai trò submodule'; "
                  "hoặc truyền --into <đường dẫn xlsx>.")
            return 2
        if not target.is_absolute():
            target = ROOT / target
    try:
        warnings = apply(spec, target, fresh=(cmd == "build"))
    except ValueError as e:
        print(f"LỖI: {e}")
        return 1
    for w in warnings:
        print(f"CẢNH BÁO: {w}")
    print(f"OK: đã ghi {target}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
